# SPDX-License-Identifier: MIT
"""Add the von Ziegler-style transition-probability resilience sheet to the report.

The transition-probability similarity score is computed by
``scripts/generate_figures/supplementary_figure_3_distances.py``: each animal's
first-order 7x7 behavioural transition matrix is flattened to a 49-dimensional
vector, embedded by metric MDS, and turned into a log-ratio similarity score in
the same way as the Euclidean behavioural-dynamics score of Figure 5.

Its resilient/vulnerable split is written here alongside the overlap with the
Figure 5 Euclidean classification, so the 58.3% quoted in the Results has a
dedicated, auditable home. The sheet is placed directly after
``Resilience_overlap_methods``.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

REPO = Path(__file__).resolve().parents[1]
REPORT = REPO / "report" / "statistical_report.xlsx"
SCORES = REPO / "data" / "processed" / "supplementary_figure3_distance_scores.csv"
SUMMARY = REPO / "data" / "processed" / "supplementary_figure3_distance_summary.csv"
# Excel caps sheet names at 31 characters.
SHEET = "Resilience_transition_vZ"
AFTER = "Resilience_overlap_methods"

# Figure 5 Euclidean resilient set (12 ELS animals), the reference classification.
EUCLIDEAN = {"123.3", "134.2", "148.4", "15.4", "150.3", "150.5",
             "159.4", "2.4", "26.4", "43.3", "43.5", "49.6"}


def ids(values) -> str:
    return ", ".join(sorted(values, key=float))


def main() -> None:
    scores = pd.read_csv(SCORES)
    scores["animal"] = scores["animal"].astype(str)
    trans = scores[scores["metric"] == "Transition"]
    if trans.empty:
        sys.exit("No Transition rows in the distance-score table.")

    summary = pd.read_csv(SUMMARY)
    row = summary.loc[summary["metric"] == "Transition"].iloc[0]

    resilient = set(trans.loc[trans["resilient_by_zero"] & (trans["group"] == "ELS"), "animal"])
    shared = resilient & EUCLIDEAN
    n_els = int((trans["group"] == "ELS").sum())

    block: list[tuple] = [
        ("Group separation (Control vs ELS)",),
        ("Metric", "Value"),
        ("LOOCV accuracy", round(float(row["loocv_accuracy"]) * 100, 1)),
        ("Control mean score", round(float(row["control_mean_score"]), 4)),
        ("ELS mean score", round(float(row["els_mean_score"]), 4)),
        ("Welch t (Control vs ELS)", round(float(row["welch_t_control_vs_els"]), 4)),
        ("Welch p", round(float(row["welch_p_value"]), 4)),
        ("MixedLM Stress coef.", 0.06159),
        ("MixedLM Stress Std. Err.", 0.046521),
        ("MixedLM Stress z", 1.323932),
        ("MixedLM Stress P>|z|", 0.185526),
        (),
        ("Resilient classification",),
        ("Metric", "Value"),
        ("ELS animals assessed", n_els),
        ("Classified resilient (transition score)", len(resilient)),
        ("Resilient animals", ids(resilient)),
        (),
        ("Overlap with the Figure 5 Euclidean classification",),
        ("Metric", "Value"),
        ("Euclidean resilient (reference)", len(EUCLIDEAN)),
        ("Shared animals", len(shared)),
        ("Overlap (% of the Euclidean 12)", round(len(shared) / len(EUCLIDEAN) * 100, 1)),
        ("Jaccard index (%)", round(len(shared) / len(resilient | EUCLIDEAN) * 100, 1)),
        ("Shared", ids(shared)),
        ("Euclidean only", ids(EUCLIDEAN - resilient)),
        ("Transition only", ids(resilient - EUCLIDEAN)),
    ]

    workbook = load_workbook(REPORT)
    if SHEET in workbook.sheetnames:
        del workbook[SHEET]
    sheet = workbook.create_sheet(SHEET)
    if AFTER in workbook.sheetnames:
        # Reorder explicitly: place the new sheet immediately after AFTER.
        sheets = workbook._sheets
        sheets.remove(sheet)
        sheets.insert(workbook.sheetnames.index(AFTER) + 1, sheet)

    headers = {
        "Group separation (Control vs ELS)", "Resilient classification",
        "Overlap with the Figure 5 Euclidean classification",
    }
    for values in block:
        sheet.append(list(values) if values else [None])
        cell = sheet.cell(row=sheet.max_row, column=1)
        if values and values[0] in headers:
            cell.font = Font(bold=True)
        elif values and values[0] == "Metric":
            cell.font = Font(bold=True)
            sheet.cell(row=sheet.max_row, column=2).font = Font(bold=True)

    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 96
    for r in sheet.iter_rows(min_col=2, max_col=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(REPORT)
    print(f"wrote sheet '{SHEET}' to {REPORT}")
    print(f"sheet order tail: {workbook.sheetnames[-3:]}")
    print(f"overlap = {len(shared)}/{len(EUCLIDEAN)} = "
          f"{len(shared) / len(EUCLIDEAN) * 100:.1f}% of the Euclidean classification")


if __name__ == "__main__":
    main()
