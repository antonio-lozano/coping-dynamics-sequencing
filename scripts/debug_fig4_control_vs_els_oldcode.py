"""Compare Fig. 4 Report.xlsx against old-code and current R/OLS models.

Scope: Control vs ELS only, n=82. Here ELS includes all ELS animals, including
the resilient animals used later in subgroup analyses.

Old-code path follows E:/#Deeplabcut_project/Code/Frequency_indexes_bout_duration_resilience.py:
  - raw 250 ms syllable table
  - assign behavior cluster
  - frequency metrics from groupby(...).size()
  - CUI from relative usage after dropping blank cluster
  - bout duration from contiguous Cluster runs
  - statsmodels MixedLM(..., groups=Animal).fit(reml=False)

The current R model used for verification is equivalent to OLS:
  metric ~ group + experiment
because there is one summary row per animal.
"""
from __future__ import annotations

from pathlib import Path
import warnings

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import entropy
import statsmodels.formula.api as smf

warnings.filterwarnings("ignore")

REPO = Path(__file__).resolve().parents[1]
RAW_250MS = REPO / "data" / "source" / "syllable_usage_per_timebin_250ms.csv"
REPORT_XLSX = Path("E:/#Deeplabcut_project/Report/Report.xlsx")

CLUSTER_MAP = {
    "Freezing": [0, 28],
    "Sniffing": [18, 20],
    "Grooming": [24],
    "Turn": [1, 3, 5, 6, 10, 15, 26, 27],
    "Locomotion": [11, 12, 14, 16, 19, 21, 25],
    "Climbing": [111],
    "Jump": [23, 29, 30, 34],
}


def load_oldcode_raw() -> pd.DataFrame:
    raw = pd.read_csv(RAW_250MS).rename(columns={"Time Bin": "Time_bin"})
    raw = raw[raw["Condition"].isin(["Control", "ELS"])].copy()
    raw["Animal"] = raw["Animal"].astype(str)
    raw["Syllable"] = pd.to_numeric(raw["Syllable"], errors="coerce").astype(int)
    raw["Experiment"] = pd.to_numeric(raw["Experiment"], errors="coerce")

    s2c: dict[int, str] = {}
    for cluster, syllables in CLUSTER_MAP.items():
        for syllable in syllables:
            s2c[int(syllable)] = cluster
    raw["Cluster"] = raw["Syllable"].map(s2c).fillna("")
    return raw


def oldcode_frequency_metrics(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    freq_df = (
        raw.groupby(["Animal", "Experiment", "Condition", "Cluster"])
        .size()
        .reset_index(name="Count")
    )
    pivot = freq_df.pivot_table(
        index=["Animal", "Condition", "Experiment"],
        columns="Cluster",
        values="Count",
        fill_value=0,
    )
    rel = pivot.div(pivot.sum(axis=1), axis=0)
    ent = rel.apply(lambda row: entropy(row[row > 0]), axis=1)
    nclus = (pivot > 0).sum(axis=1)
    even = ent / np.log(nclus)
    simp = rel.apply(lambda row: 1 - np.sum(row**2), axis=1)
    metrics = pd.DataFrame(
        {"Shannon": ent, "Evenness": even, "Simpson": simp}
    ).reset_index()
    metrics["group"] = metrics["Condition"]
    return metrics, freq_df


def oldcode_cui(freq_df: pd.DataFrame, raw: pd.DataFrame) -> pd.DataFrame:
    # Merge original Condition, then compute relative usage per animal.
    if "Condition" in freq_df.columns:
        freq_use = freq_df.drop(columns=["Condition"]).copy()
    else:
        freq_use = freq_df.copy()
    merged = freq_use.merge(
        raw[["Animal", "Condition"]].drop_duplicates(),
        on="Animal",
        how="left",
    )
    total = (
        merged.groupby(["Animal", "Condition", "Experiment"])["Count"]
        .sum()
        .reset_index(name="Total")
    )
    usage = merged.merge(total, on=["Animal", "Condition", "Experiment"])
    usage["RelativeUsage"] = usage["Count"] / usage["Total"]
    usage = usage.dropna(subset=["Cluster"])
    usage = usage[usage["Cluster"].astype(str).str.strip() != ""].copy()

    def compute_cui(subdf: pd.DataFrame) -> float:
        sorted_df = subdf.sort_values("RelativeUsage", ascending=False)
        cumulative = sorted_df["RelativeUsage"].cumsum().to_numpy()
        n = len(cumulative)
        baseline = (n + 1) / (2 * n)
        return float((cumulative.mean() - baseline) / (1 - baseline))

    out = (
        usage.groupby(["Animal", "Experiment", "Condition"])
        .apply(compute_cui)
        .reset_index(name="CUI")
    )
    out["group"] = out["Condition"]
    return out


def oldcode_cluster_bouts(raw: pd.DataFrame) -> pd.DataFrame:
    def compute_cluster_bout_duration(subdf: pd.DataFrame) -> pd.Series:
        subdf = subdf.sort_values("Time_bin").copy()
        subdf["Cluster_change"] = (subdf["Cluster"] != subdf["Cluster"].shift(1)).astype(int)
        subdf["Bout"] = subdf["Cluster_change"].cumsum()
        bouts = subdf.groupby(["Bout", "Cluster"]).size() * 0.25
        return bouts.groupby(level=1).mean()

    out = raw.groupby(["Animal", "Experiment"]).apply(compute_cluster_bout_duration).reset_index()
    out.columns = ["Animal", "Experiment", "Cluster", "MeanBoutDuration"]
    out = out.dropna(subset=["MeanBoutDuration"]).reset_index(drop=True)
    out = out.merge(
        raw[["Animal", "Condition", "Experiment"]].drop_duplicates(),
        on=["Animal", "Experiment"],
        how="left",
    )
    out["group"] = out["Condition"]
    return out


def fit_old_mixedlm(df: pd.DataFrame, y: str) -> tuple[float, float, float, float, int]:
    sub = df.dropna(subset=[y]).copy()
    sub["group"] = pd.Categorical(sub["group"], categories=["Control", "ELS"])
    model = smf.mixedlm(f"{y} ~ group + Experiment", sub, groups=sub["Animal"]).fit(reml=False)
    key = "group[T.ELS]"
    beta = float(model.params[key])
    se = float(model.bse[key])
    z = float(model.tvalues[key])
    p = float(model.pvalues[key])
    return beta, se, z, p, int(model.nobs)


def fit_current_ols(df: pd.DataFrame, y: str) -> tuple[float, float, float, float, int]:
    sub = df.dropna(subset=[y]).copy()
    sub["group"] = pd.Categorical(sub["group"], categories=["Control", "ELS"])
    model = smf.ols(f"{y} ~ group + Experiment", sub).fit()
    key = "group[T.ELS]"
    beta = float(model.params[key])
    se = float(model.bse[key])
    t = float(model.tvalues[key])
    p = float(model.pvalues[key])
    return beta, se, t, p, int(model.nobs)


def report_value(sheet: str, start_col: int) -> tuple[float, float, float, float, int]:
    wb = load_workbook(REPORT_XLSX, read_only=True, data_only=True)
    ws = wb[sheet]
    return (
        float(ws.cell(4, start_col + 1).value),
        float(ws.cell(4, start_col + 2).value),
        float(ws.cell(4, start_col + 3).value),
        float(ws.cell(4, start_col + 4).value),
        int(ws.cell(6, start_col + 10).value),
    )


def fmt(vals: tuple[float, float, float, float, int]) -> str:
    return f"b={vals[0]:+.6f} se={vals[1]:.6f} z/t={vals[2]:+.6f} p={vals[3]:.6g} n={vals[4]}"


def main() -> None:
    raw = load_oldcode_raw()
    metrics, freq_df = oldcode_frequency_metrics(raw)
    cui = oldcode_cui(freq_df, raw)
    bouts = oldcode_cluster_bouts(raw)

    print("Input")
    print("rows", len(raw), "animals", raw["Animal"].nunique())
    print("groups", raw.groupby("Condition")["Animal"].nunique().to_dict())
    print("experiments", raw.groupby("Experiment")["Animal"].nunique().to_dict())

    checks = [
        ("Fig4E Simpson", metrics, "Simpson", "Frequency_metrics", 1),
        ("Fig4F Shannon", metrics, "Shannon", "Frequency_metrics", 14),
        ("Fig4G Evenness", metrics, "Evenness", "Frequency_metrics", 27),
        ("Fig4H-I CUI", cui, "CUI", "Frequency_metrics", 41),
        ("Fig4K Freezing bout", bouts[bouts["Cluster"] == "Freezing"], "MeanBoutDuration", "Bout_duration", 14),
        ("Fig4L Sniffing bout", bouts[bouts["Cluster"] == "Sniffing"], "MeanBoutDuration", "Bout_duration", 27),
        ("Fig4N Turn bout", bouts[bouts["Cluster"] == "Turn"], "MeanBoutDuration", "Bout_duration", 54),
    ]

    print("\nComparison: Report.xlsx vs old-code collapsed MixedLM vs current R/OLS equivalent")
    for label, df, y, sheet, col in checks:
        rep = report_value(sheet, col)
        old = fit_old_mixedlm(df, y)
        ols = fit_current_ols(df, y)
        print(f"\n{label}")
        print("  Report.xlsx:        ", fmt(rep))
        print("  Old-code MixedLM:   ", fmt(old))
        print("  Current R/OLS equiv:", fmt(ols))


if __name__ == "__main__":
    main()
