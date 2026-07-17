"""Inspect Report.xlsx sheets relevant to Figure 4."""
import openpyxl

path = r"E:\#Deeplabcut_project\Report\Report.xlsx"
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print("All sheets:")
for name in wb.sheetnames:
    print(" ", name)

# Look for sheets likely related to Figure 4 (diversity, bout, transition, frequency)
keywords = ["freq", "bout", "transit", "diversity", "shannon", "simpson", "lz", "recur", "determin", "markov", "fig4", "figure4"]
print("\nPotentially relevant sheets:")
for name in wb.sheetnames:
    if any(k in name.lower() for k in keywords):
        print(f"\n=== {name} ===")
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        for row in rows[:30]:
            if any(cell is not None for cell in row):
                print(row)
wb.close()
