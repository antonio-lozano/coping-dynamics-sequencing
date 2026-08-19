#!/usr/bin/env python
"""Validate the publication repository from a clean clone."""
from __future__ import annotations

import csv
import hashlib
import re
import sys
from pathlib import Path

import openpyxl
import imageio_ffmpeg
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

# The manifest writer owns the definition of what counts as a publication
# artifact and which files a clean clone contains. Importing it here keeps one
# definition rather than two that can drift apart unnoticed.
sys.path.insert(0, str(Path(__file__).resolve().parent))

from update_manifest import clean_clone_paths, iter_artifacts, sha256  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]

REQUIRED_FILES = [
    "README.md",
    "LICENSE",
    "requirements.txt",
    "environment.yml",
    "CITATION.cff",
    ".gitattributes",
    ".gitignore",
    "matplotlibrc",
    ".github/workflows/reproducibility.yml",
    "MANIFEST.csv",
    "REPRODUCIBILITY.md",
    "DATA_AVAILABILITY.md",
    "CODE_AVAILABILITY.md",
    "docs/data_dictionary.md",
    "docs/figure4_coping_provenance.md",
    "docs/figure_structure.md",
    "docs/behavior_classifier.md",
    "docs/supplementary_media.md",
    "docs/workbook_match_audit.md",
    "data/raw/animal_groups.csv",
    "data/raw/behavioral_flexibility_scores.xlsx",
    "data/raw/freezing_overlap_by_group.csv",
    "data/raw/freezing_predictions_index.csv",
    "data/raw/freezing_predictions_light.csv.gz",
    "data/raw/moseq_syllables_per_frame.csv.gz",
    "data/raw/simba_validation_manual_vs_automatic.csv",
    "data/raw/syllable_classification_metrics.csv",
    "data/raw/syllable_usage_per_timebin_30s.csv",
    "data/raw/syllable_usage_per_timebin_250ms.csv",
    "data/raw/updated_results.pkl.gz",
    "data/raw/manuscript_tables/statistical_report/provenance.md",
    "data/raw/manuscript_tables/statistical_report/reported_cells.jsonl",
    "data/raw/manuscript_tables/statistical_report/reported_layout.json",
    "data/processed/cluster_frequency_per_animal.csv",
    "data/processed/cluster_timecourse_per_animal.csv",
    "data/processed/s0s28_timecourse_per_animal.csv",
    "data/processed/supplementary_figure1_tracking_clusters.csv",
    "data/processed/supplementary_figure1_time_summary.csv",
    "data/processed/supplementary_figure3_distance_scores.csv",
    "data/processed/supplementary_figure3_distance_summary.csv",
    "data/processed/supplementary_figure3_threshold_audit.csv",
    "data/processed/tracking_exclusions_per_animal.csv",
    "data/processed/transition_metrics_per_animal.csv",
    "data/processed/figure5_dynamics_scores.csv",
    "data/processed/figure5_resilience_threshold_audit.csv",
    "figure_source_data/figure2.csv",
    "figure_source_data/figure3.csv",
    "figure_source_data/figure4.csv",
    "figure_source_data/figure5.csv",
    "figure_source_data/figure6.csv",
    "figure_source_data/figure7.csv",
    "figure_source_data/figure7_classifier_accuracy.csv",
    "figure_source_data/figure7_classifier_confusion_matrix.csv",
    "figure_source_data/figure7_classifier_global_shap.csv",
    "figure_source_data/figure7_full_session_auc.csv",
    "figure_source_data/figure7_individual_timecourse_auc.csv",
    "figure_source_data/figure7_prediction_onset.csv",
    "figure_source_data/figure7_predictor_catalog.csv",
    "figure_source_data/figure7_predictor_timecourse.csv",
    "figure_source_data/figure7_shapley_contributions.csv",
    "figure_source_data/figure7_shapley_per_animal.csv",
    "figure_source_data/supplementary_figure5_shap_summary.csv",
    "statistics/stats_figure3A_GEE.csv",
    "statistics/stats_figure3_overtime.csv",
    "statistics/stats_figure2_ground_truth_MixedLM.csv",
    "statistics/stats_figure2h_syllables_MixedLM.csv",
    "statistics/stats_figure5_dynamics.csv",
    "statistics/stats_figure5_timecourse_MixedLM.csv",
    "statistics/stats_figure4_bouts_MixedLM.csv",
    "statistics/stats_figure4_diversity_MixedLM.csv",
    "statistics/stats_figure4_transition_MixedLM.csv",
    "statistics/fig4_bout_cluster_per_animal.csv",
    "statistics/fig4_bout_overall_per_animal.csv",
    "statistics/fig4_diversity_per_animal.csv",
    "statistics/fig4_transition_per_animal.csv",
    "statistics/fig5c_frequency_gee.csv",
    "statistics/fig6_bout_resilience_stats.csv",
    "statistics/fig6_diversity_resilience_stats.csv",
    "statistics/fig6_transition_resilience_stats.csv",
    "statistics/figure7_cross_cohort_auc.csv",
    "statistics/figure7_full_session_auc.csv",
    "statistics/figure7_prediction_permutation.csv",
    "classifier/figure7_behavior_classifier.joblib",
    "supplementary_media/Supplementary_Video_1_MoSeq_syllable_atlas.mp4",
    "supplementary_media/Supplementary_Video_1_source_index.csv",
    "report/raw_data.xlsx",
    "report/statistical_report.xlsx",
    "report/SIGuide.docx",
    "report/Sanguino-Gómez_coping_strategies_Nature_Neuroscience_highlighted.docx",
    "scripts/run_all.py",
    "scripts/run_all_figures.py",
    "scripts/update_manifest.py",
    "scripts/check_manuscript.py",
    "scripts/build_raw_data_workbook.py",
    "scripts/build_statistical_report.py",
    "scripts/save_deterministic.py",
    "scripts/compare_workbook_data.py",
    "scripts/import_legacy_statistical_report.py",
    "scripts/derive_tables/freezing_predictions_light.py",
    "scripts/derive_tables/cluster_tables.py",
    "scripts/derive_tables/tracking_exclusions.py",
    "scripts/derive_tables/fig6_resilience_stats.py",
    "scripts/generate_figures/figure_7_resilience_prediction.py",
    "scripts/generate_figures/supplementary_figure_5_classifier_shap.py",
    "scripts/generate_supplementary_video_1.py",
    "scripts/migrations/import_legacy_shap_summary.py",
    "scripts/migrations/import_legacy_shap_values.py",
]

FIGURE_STEMS = [
    "figure1",
    "figure2",
    "figure3",
    "figure4",
    "figure5",
    "figure6",
    "figure7",
    "supplementary_figure1",
    "supplementary_figure2",
    "supplementary_figure3",
    "supplementary_figure4",
    "supplementary_figure5",
]
FIGURE_EXTS = (".pdf", ".svg", ".png")

FORBIDDEN_TOP_LEVEL = {"results", "dataset", "external"}
FORBIDDEN_PATHS = {"data/" + "derived"}
TEXT_SUFFIXES = {".md", ".py", ".txt", ".json", ".yml", ".yaml", ".cff", ".toml"}
LOCAL_PATH_PATTERN = re.compile(r"(?i)\b[a-z]:\\(?:users|downloads|jen|antonio|big_computer)|/mnt/[a-z]/")
SKIP_DIRS = {".git", ".venv", "__pycache__", ".mypy_cache", ".pytest_cache", ".uv-cache"}
# Reading a file back to explain a hash mismatch is only worth it while it fits
# comfortably in memory; past that the plain mismatch is the more useful report.
LINE_ENDING_HINT_MAX_BYTES = 64 * 1024 * 1024
FORBIDDEN_TEXT_TERMS = (
    "Cl" + "aude",
    "Co" + "dex",
    "Chat" + "GPT",
    "Open" + "AI",
    "AI" + "-generated",
    "AI" + " generated",
)


def fail(message: str, errors: list[str]) -> None:
    errors.append(message)


def path_exists_exact_case(rel: str) -> bool:
    current = ROOT
    for part in Path(rel).parts:
        if not current.is_dir():
            return False
        entries = {child.name: child for child in current.iterdir()}
        if part not in entries:
            return False
        current = entries[part]
    return True


def iter_files() -> list[Path]:
    """Every file a clean clone would contain.

    Files Git ignores are left out, so a working copy holding local scratch
    reports exactly what a continuous integration runner reports. Untracked files
    that are *not* ignored stay in, because those are the ones that would reach
    the repository unnoticed on the next commit.
    """
    included = clean_clone_paths()
    candidates = (ROOT / rel for rel in included) if included is not None else ROOT.rglob("*")
    return [
        path
        for path in candidates
        if not SKIP_DIRS.intersection(path.parts) and path.is_file()
    ]


def check_required(errors: list[str]) -> None:
    for rel in REQUIRED_FILES:
        path = ROOT / rel
        if not path.is_file():
            fail(f"missing required file: {rel}", errors)
        elif path.stat().st_size == 0:
            fail(f"empty required file: {rel}", errors)

    for stem in FIGURE_STEMS:
        for ext in FIGURE_EXTS:
            rel = f"figures/{stem}{ext}"
            path = ROOT / rel
            if not path.is_file():
                fail(f"missing figure export: {rel}", errors)

    pred_dir = ROOT / "data/raw/freezing_predictions"
    predictions = sorted(pred_dir.glob("*.csv")) if pred_dir.exists() else []
    if len(predictions) != 98:
        fail(f"expected 98 freezing prediction CSVs, found {len(predictions)}", errors)


def check_layout(files: list[Path], errors: list[str]) -> None:
    for name in FORBIDDEN_TOP_LEVEL:
        if (ROOT / name).exists():
            fail(f"forbidden top-level path exists: {name}", errors)
    for rel in FORBIDDEN_PATHS:
        if path_exists_exact_case(rel):
            fail(f"forbidden retired path exists: {rel}", errors)

    for path in files:
        rel = path.relative_to(ROOT).as_posix()
        lower = rel.lower()
        if path.name.lower().startswith("readme") and rel != "README.md":
            fail(f"subfolder README found: {rel}", errors)
        if (
            lower.startswith("data/raw/")
            and path.suffix.lower() in {".pdf", ".png", ".svg"}
        ):
            fail(f"figure artifact found in raw data: {rel}", errors)
        if lower.endswith(("_test.xlsx", "_rebuilt.xlsx")) or "smoke" in lower:
            fail(f"scratch artifact found: {rel}", errors)
        if ("no" + "seed") in lower or ("no" + "-seed") in lower:
            fail(f"analysis-decision scratch name found: {rel}", errors)


def check_text_clean(files: list[Path], errors: list[str]) -> None:
    for path in files:
        if path.suffix.lower() not in TEXT_SUFFIXES:
            continue
        rel = path.relative_to(ROOT).as_posix()
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            continue
        for match in LOCAL_PATH_PATTERN.finditer(text):
            fail(f"local absolute path in {rel}: {match.group(0)}", errors)
        for term in FORBIDDEN_TEXT_TERMS:
            if term in text:
                fail(f"tool-specific text trace in {rel}", errors)


def line_ending_hint(path: Path, expected_sha: str) -> str:
    """Name line endings as the cause when the bytes differ but the content does not.

    A tree checked out with the platform's own line endings hashes differently
    from one checked out with LF. Without this, that surfaces as an unexplained
    hash mismatch on every text artifact at once, which reads like data loss.
    """
    if not expected_sha or path.stat().st_size > LINE_ENDING_HINT_MAX_BYTES:
        return ""
    as_lf = path.read_bytes().replace(b"\r\n", b"\n")
    variants = (as_lf, as_lf.replace(b"\n", b"\r\n"))
    if any(hashlib.sha256(variant).hexdigest() == expected_sha for variant in variants):
        return " (same content, different line endings; see .gitattributes)"
    return ""


def check_manifest(errors: list[str]) -> None:
    manifest = ROOT / "MANIFEST.csv"
    if not manifest.exists():
        fail("missing MANIFEST.csv", errors)
        return

    with manifest.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        fail("MANIFEST.csv has no artifact rows", errors)
        return

    recorded = {row.get("path", "") for row in rows}
    if len(recorded) != len(rows):
        fail("MANIFEST.csv lists the same path more than once", errors)

    for row in rows:
        rel = row.get("path", "")
        path = ROOT / rel
        if not path.is_file():
            fail(f"manifest path missing: {rel}", errors)
            continue

        expected_sha = row.get("sha256", "")
        try:
            expected_size = int(row.get("bytes", ""))
        except (TypeError, ValueError):
            fail(f"manifest row has a non-numeric size: {rel}", errors)
            continue

        actual_size = path.stat().st_size
        if expected_size != actual_size:
            hint = line_ending_hint(path, expected_sha)
            fail(f"manifest size mismatch: {rel} (recorded {expected_size}, found {actual_size}){hint}", errors)
        if expected_sha != sha256(path):
            fail(f"manifest sha256 mismatch: {rel}{line_ending_hint(path, expected_sha)}", errors)

    # The reverse direction. Without it, an artifact added to one of the covered
    # directories and never hashed passes silently, and the manifest stops being
    # a complete record of what the repository ships.
    for path in iter_artifacts():
        rel = path.relative_to(ROOT).as_posix()
        if rel not in recorded:
            fail(f"artifact on disk but absent from MANIFEST.csv: {rel}", errors)


def check_supplementary_video(errors: list[str]) -> None:
    video = ROOT / "supplementary_media/Supplementary_Video_1_MoSeq_syllable_atlas.mp4"
    index = ROOT / "supplementary_media/Supplementary_Video_1_source_index.csv"
    if not video.is_file() or not index.is_file():
        return

    if video.stat().st_size > 30 * 1024 * 1024:
        fail("Supplementary Video 1 exceeds 30 MB", errors)
    reader = imageio_ffmpeg.read_frames(str(video), pix_fmt="rgb24")
    try:
        metadata = next(reader)
    finally:
        reader.close()
    if metadata.get("codec") != "h264":
        fail(f"Supplementary Video 1 codec is {metadata.get('codec')}, not h264", errors)
    if not str(metadata.get("pix_fmt", "")).startswith("yuv420p"):
        fail(f"Supplementary Video 1 pixel format is {metadata.get('pix_fmt')}", errors)
    if tuple(metadata.get("size", ())) != (960, 540):
        fail(f"Supplementary Video 1 frame size is {metadata.get('size')}", errors)
    if float(metadata.get("fps", 0)) != 25.0:
        fail(f"Supplementary Video 1 frame rate is {metadata.get('fps')}", errors)

    with index.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    labels = [int(row["syllable"]) for row in rows]
    expected = set(range(35)) | {111}
    if len(rows) != 36 or set(labels) != expected or len(labels) != len(set(labels)):
        fail("Supplementary Video 1 source index does not uniquely cover 0-34 and 111", errors)
    if any(len(row.get("source_clip_sha256", "")) != 64 for row in rows):
        fail("Supplementary Video 1 source index contains an invalid SHA-256", errors)


def check_submission_manuscript(errors: list[str]) -> None:
    path = ROOT / "report/Sanguino-Gómez_coping_strategies_Nature_Neuroscience_highlighted.docx"
    if not path.is_file():
        return
    document = Document(path)
    paragraphs = document.paragraphs
    abstract_heading = next(
        (index for index, paragraph in enumerate(paragraphs) if paragraph.text == "Abstract (150 words)"),
        None,
    )
    if abstract_heading is None or len(paragraphs[abstract_heading + 1].text.split()) != 150:
        fail("Nature Neuroscience abstract is not exactly 150 space-delimited words", errors)

    video_legend = next(
        (paragraph for paragraph in paragraphs if paragraph.text.startswith("The video presents")),
        None,
    )
    if video_legend is None or len(video_legend.text.split()) > 100:
        fail("Supplementary Video 1 legend is missing or exceeds 100 words", errors)

    required_additions = (
        "Representative pose-overlaid examples and canonical skeleton trajectories",
        "For the audiovisual atlas, three 20-frame representative occurrences",
        "Supplementary Video 1. Representative MoSeq syllables",
    )
    for addition in required_additions:
        paragraph = next((item for item in paragraphs if addition in item.text), None)
        if paragraph is None:
            fail(f"submission manuscript is missing addition: {addition}", errors)
            continue
        highlighted = "".join(
            run.text
            for run in paragraph.runs
            if run.font.highlight_color == WD_COLOR_INDEX.YELLOW
        )
        if addition not in highlighted:
            fail(f"submission manuscript addition is not precisely highlighted: {addition}", errors)


def check_workbook_figure_structure(errors: list[str]) -> None:
    """Verify that Figure 7 and supplementary sheets follow the panel map."""
    expectations = {
        "report/raw_data.xlsx": [
            "Fig.7A_Classifier_accuracy",
            "Fig.7B_Classifier_SHAP",
            "Fig.7C_Confusion_matrix",
            "Fig.7D-E_Prediction_AUC",
            "Fig.7F_SHAP_contributions",
            "Fig.7G-H_Predictor_catalog",
            "Fig.7I_Prediction_onset",
            "Fig.7J_Family_timecourse",
            "Fig.7K-N_Metric_timecourse",
            "Suppl.Fig.1A-D",
            "Suppl.Fig.3A-K",
            "Suppl.Fig.5A-H_SHAP",
        ],
        "report/statistical_report.xlsx": [
            "Fig.7D-E_Prediction",
            "Fig.7F_SHAP",
            "Fig.7I_Prediction_onset",
            "Suppl.Fig.3A-K",
        ],
    }
    for rel, expected in expectations.items():
        path = ROOT / rel
        if not path.is_file():
            continue
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        missing = [name for name in expected if name not in workbook.sheetnames]
        if missing:
            fail(f"{rel} is missing canonical figure sheets: {missing}", errors)
            continue
        positions = [workbook.sheetnames.index(name) for name in expected]
        if positions != sorted(positions):
            fail(f"{rel} figure sheets are not in canonical panel order", errors)

    raw = ROOT / "report/raw_data.xlsx"
    if raw.is_file():
        workbook = openpyxl.load_workbook(raw, read_only=True, data_only=True)
        for sheet_name in [name for name in workbook.sheetnames if name.startswith("Fig.7")]:
            sheet = workbook[sheet_name]
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=3, max_row=3))]
            forbidden = {"Label", "Animal_Label", "Panel"}.intersection(headers)
            if forbidden:
                fail(f"{sheet_name} contains helper columns: {sorted(forbidden)}", errors)
            stale = {
                cell.value
                for row in sheet.iter_rows()
                for cell in row
                if cell.value in {"Exp1", "Exp3"}
            }
            if stale:
                fail(f"{sheet_name} contains abbreviated dataset names: {sorted(stale)}", errors)

    stats = ROOT / "report/statistical_report.xlsx"
    if stats.is_file():
        workbook = openpyxl.load_workbook(stats, read_only=True, data_only=True)
        raw_only = {
            "Fig.7A_Classifier_accuracy",
            "Fig.7B_Classifier_SHAP",
            "Fig.7C_Confusion_matrix",
            "Fig.7G-H_Predictors",
            "Fig.7J_Family_timecourse",
            "Fig.7K-N_Metric_timecourse",
        }
        duplicated = sorted(raw_only.intersection(workbook.sheetnames))
        if duplicated:
            fail(f"statistical report duplicates raw Figure 7 sheets: {duplicated}", errors)


def main() -> int:
    errors: list[str] = []
    files = iter_files()
    check_required(errors)
    check_layout(files, errors)
    check_text_clean(files, errors)
    check_manifest(errors)
    check_supplementary_video(errors)
    check_submission_manuscript(errors)
    check_workbook_figure_structure(errors)

    if errors:
        print("Reproducibility checks failed:")
        for err in errors:
            print(f"  - {err}")
        return 1

    print("Reproducibility checks passed.")
    print("Verified required inputs, processed data, source data, figures, reports, classifier artifact, and hashes.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
