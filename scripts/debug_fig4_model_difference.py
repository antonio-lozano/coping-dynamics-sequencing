from __future__ import annotations

import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import statsmodels.formula.api as smf
from openpyxl import load_workbook

warnings.filterwarnings("ignore")

RAW_XLSX = Path("D:/coping-dynamics-sequencing/data/Raw_data.xlsx")
REPORT_XLSX = Path("E:/#Deeplabcut_project/Report/Report.xlsx")


def read_raw_sheet(sheet: str) -> pd.DataFrame:
    raw = pd.read_excel(RAW_XLSX, sheet_name=sheet, header=None)
    header_idx = None
    for idx, row in raw.iterrows():
        vals = set(row.astype(str))
        if {"animal_id", "group", "experiment"} & vals:
            header_idx = idx
            break
    if header_idx is None:
        raise RuntimeError(f"No header in {sheet}")
    headers = raw.iloc[header_idx].tolist()
    df = raw.iloc[header_idx + 1 :].copy()
    df.columns = headers
    df = df.loc[:, [c for c in df.columns if pd.notna(c)]]
    return df.dropna(how="all").infer_objects()


def report_value(sheet: str, start_col: int) -> tuple[float, float, float, float]:
    wb = load_workbook(REPORT_XLSX, data_only=True, read_only=True)
    ws = wb[sheet]
    return (
        float(ws.cell(4, start_col + 1).value),
        float(ws.cell(4, start_col + 2).value),
        float(ws.cell(4, start_col + 3).value),
        float(ws.cell(4, start_col + 4).value),
    )


def fit_all(df: pd.DataFrame, y: str, label: str, report: tuple[float, float, float, float]) -> None:
    sub = df.dropna(subset=[y]).copy()
    sub[y] = pd.to_numeric(sub[y])
    sub["animal_id"] = sub["animal_id"].astype(str)
    sub["group"] = pd.Categorical(sub["group"], categories=["Control", "ELS"])
    sub["experiment_num"] = pd.to_numeric(sub["experiment"])
    sub["experiment_cat"] = pd.Categorical(sub["experiment"].astype(str))
    print(f"\n== {label} ==")
    print("Report beta/se/z/p:", tuple(round(x, 6) for x in report))
    print("n", len(sub), "groups", sub["group"].value_counts().to_dict(), "experiments", sub["experiment"].value_counts().to_dict())

    variants = [
        ("OLS numeric experiment", lambda: smf.ols(f"{y} ~ group + experiment_num", sub).fit(), "group[T.ELS]"),
        ("OLS categorical experiment", lambda: smf.ols(f"{y} ~ group + experiment_cat", sub).fit(), "group[T.ELS]"),
        ("OLS group only", lambda: smf.ols(f"{y} ~ group", sub).fit(), "group[T.ELS]"),
        ("MixedLM formula numeric experiment", lambda: smf.mixedlm(f"{y} ~ group + experiment_num", sub, groups=sub["animal_id"]).fit(reml=False), "group[T.ELS]"),
        ("MixedLM formula categorical experiment", lambda: smf.mixedlm(f"{y} ~ group + experiment_cat", sub, groups=sub["animal_id"]).fit(reml=False), "group[T.ELS]"),
        ("MixedLM formula group only", lambda: smf.mixedlm(f"{y} ~ group", sub, groups=sub["animal_id"]).fit(reml=False), "group[T.ELS]"),
    ]
    for name, maker, key in variants:
        try:
            m = maker()
            beta = float(m.params[key])
            se = float(m.bse[key])
            stat = beta / se
            p = float(m.pvalues[key])
            print(f"{name:38s} beta={beta:+.6f} se={se:.6f} z/t={stat:+.6f} p={p:.6g}")
        except Exception as exc:
            print(f"{name:38s} failed: {type(exc).__name__}: {exc}")


def main() -> None:
    div = read_raw_sheet("Fig4_frequency_metrics")
    bouts = read_raw_sheet("Fig4_bout_duration")
    fits = [
        (div, "simpson_index", "Simpson", report_value("Frequency_metrics", 1)),
        (div, "cumulative_usage_index", "CUI", report_value("Frequency_metrics", 41)),
        (bouts[bouts["cluster"] == "Freezing"], "bout_duration_seconds", "Freezing bout", report_value("Bout_duration", 14)),
        (bouts[bouts["cluster"] == "Sniffing"], "bout_duration_seconds", "Sniffing bout", report_value("Bout_duration", 27)),
        (bouts[bouts["cluster"] == "Turn"], "bout_duration_seconds", "Turn bout", report_value("Bout_duration", 54)),
    ]
    for args in fits:
        fit_all(*args)


if __name__ == "__main__":
    main()
