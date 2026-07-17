"""Quick check of the Fig.4 sheet in Statistical_report.xlsx."""
import openpyxl

wb = openpyxl.load_workbook(
    "d:/coping-dynamics-sequencing/results/statistical_reports/Statistical_report.xlsx",
    read_only=True, data_only=True
)
print("Sheets:", wb.sheetnames)
ws = wb["Fig.4_Diversity_Dynamics"]
rows = list(ws.iter_rows(values_only=True))
print(f"\nTotal rows: {len(rows)}")
for i, row in enumerate(rows[:80]):
    if any(c is not None for c in row):
        # Print only first 5 cols
        print(f"  row {i+1}: {row[:5]}")
wb.close()
