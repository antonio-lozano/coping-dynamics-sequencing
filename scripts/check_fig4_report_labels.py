from __future__ import annotations

from openpyxl import load_workbook

REPORT = "E:/#Deeplabcut_project/Report/Report.xlsx"

BLOCKS = [
    ("Frequency_metrics", "Simpson", 1),
    ("Frequency_metrics", "Shannon", 14),
    ("Frequency_metrics", "Evenness", 27),
    ("Frequency_metrics", "CUI", 41),
    ("Bout_duration", "Overall", 1),
    ("Bout_duration", "Freezing", 14),
    ("Bout_duration", "Sniffing", 27),
    ("Bout_duration", "Turn", 54),
]


def main() -> None:
    wb = load_workbook(REPORT, read_only=True, data_only=True)
    for sheet, label, col in BLOCKS:
        ws = wb[sheet]
        print(f"\n[{sheet} - {label}]")
        for r in range(1, 8):
            vals = [ws.cell(r, c).value for c in range(col, col + 12)]
            print(r, vals)


if __name__ == "__main__":
    main()
