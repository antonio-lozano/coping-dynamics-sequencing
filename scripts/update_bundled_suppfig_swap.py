#!/usr/bin/env python
# SPDX-License-Identifier: MIT
"""Swap Supplementary Figures 4 and 5 in the bundled manuscript copies.

The supplementary figures were cited out of numerical order: the SHAP figure
(numbered 5) is first mentioned in the Results, before the climbing-validation
figure (numbered 4) is first mentioned in the Methods. Journals number
supplementary items by order of first mention, so the two swap numbers:

    Supplementary Fig. 4   class-specific classifier SHAP beeswarms
    Supplementary Fig. 5   convex-hull floor-overlap climbing validation

The repository side of the swap (figures/, figure_source_data/, the generator
script, workbook sheet names, docs) is handled in the same change set; this
script renumbers the three citations and the two legend headings inside the
bundled .docx copies and swaps the two legend paragraphs so the legends list
stays in numerical order. Replacements are highlighted bright green.

Like the other update_bundled_* scripts this is a one-time migration: the
before strings match the bundled copies as they are today.

Run:
    python scripts/update_bundled_suppfig_swap.py
"""

from __future__ import annotations

import docx
from update_bundled_manuscripts import TARGETS, applied, problems, replace

EN = "–"

EDITS: list[tuple[str, str, str]] = [
    (
        "(Supplementary Fig. 5A" + EN + "H)",
        "(Supplementary Fig. 4A" + EN + "H)",
        "Results citation of the SHAP panels",
    ),
    ("(Supplementary Fig. 5)", "(Supplementary Fig. 4)", "Results citation of the SHAP beeswarms"),
    (
        "(Supplementary Figure 4)",
        "(Supplementary Figure 5)",
        "Methods citation of the climbing validation",
    ),
    (
        "Supplementary Fig. 4. Convex-hull floor-overlap metric",
        "Supplementary Fig. 5. Convex-hull floor-overlap metric",
        "climbing legend heading",
    ),
    (
        "Supplementary Fig. 5. Class-specific SHAP beeswarm plots",
        "Supplementary Fig. 4. Class-specific SHAP beeswarm plots",
        "SHAP legend heading",
    ),
]


def main() -> int:
    for target in TARGETS:
        doc = docx.Document(str(target))
        done_before = len(applied)
        for old, new, label in EDITS:
            hit = False
            for paragraph in doc.paragraphs:
                if replace(paragraph, old, new, label=f"{target.name}: {label}"):
                    hit = True
                    break
            if not hit:
                problems.append(f"NOT FOUND in {target.name}: [{label}]")

        # The legends list must stay in numerical order, so the SHAP legend
        # (now Supplementary Fig. 4) moves in front of the climbing legend
        # (now Supplementary Fig. 5).
        climb = shap = None
        for paragraph in doc.paragraphs:
            if paragraph.text.startswith("Supplementary Fig. 5. Convex-hull"):
                climb = paragraph
            elif paragraph.text.startswith("Supplementary Fig. 4. Class-specific"):
                shap = paragraph
        if climb is None or shap is None:
            problems.append(f"{target.name}: could not locate both legend paragraphs to reorder")
        else:
            climb._p.addprevious(shap._p)
            applied.append(f"{target.name}: legends reordered")

        doc.save(str(target))
        print(f"{target.name}: {len(applied) - done_before}/{len(EDITS) + 1} edits applied")

    for line in problems:
        print(line)
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
