"""Check Fig.4 sheet dimensions and a few key cells."""
import openpyxl
wb = openpyxl.load_workbook(
    "d:/coping-dynamics-sequencing/results/statistical_reports/Statistical_report.xlsx",
    data_only=True
)
ws = wb["Fig.4_Diversity_Dynamics"]
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
for r in range(1, min(ws.max_row+1, 120)):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 6)]
    if any(v is not None for v in row_vals):
        print(f"  r{r}: {row_vals}")
wb.close()
