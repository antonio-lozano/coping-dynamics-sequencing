# SPDX-License-Identifier: MIT
"""Point-by-point audit of the manuscript against the repository.

Checks, in order:
  1. figure and table citation order (numbered in order of first mention);
  2. panel letters used in the text vs the panels each figure actually draws;
  3. legend panel letters vs the same;
  4. reference citation order vs the reference list;
  5. figure/legend inventory vs the artifacts the repo ships.

Writes a markdown report; every finding is a line the authors can act on.
"""

from __future__ import annotations

import re
import sys
from collections import OrderedDict
from pathlib import Path

import docx

REPO = Path(__file__).resolve().parents[1]
# The manuscript lives outside the repository, so its path is supplied at run
# time: `python scripts/audit_manuscript_vs_repo.py PATH/TO/manuscript.docx`.
if len(sys.argv) > 1:
    DOC = Path(sys.argv[1])
else:
    _found = sorted((Path.home() / "Downloads").glob("*coping_strategies*.docx"))
    if not _found:
        sys.exit("Pass the manuscript .docx path as the first argument.")
    DOC = _found[0]
OUT = REPO / "statistics" / "MANUSCRIPT_REPO_AUDIT.md"

# Panels each figure actually contains, read off the generators / rendered files.
FIGURE_PANELS = {
    "1": list("AB"),
    "2": list("ABCDEFGH"),
    "3": list("ABCDEFGH"),
    "4": list("ABCDEFGHIJKLMNOPQRSTUVW"),
    "5": list("ABCDEFGHIJ"),
    "6": list("ABCDEFGHIJKLMNOPQRSTUVWXYZ"),
    "7": list("ABCDEFGHIJKLMN"),
}
SUPP_PANELS = {
    "1": list("ABCD"),
    "2": [],
    "3": list("ABCDEFGH"),
    "4": [],
    "5": list("ABCDEFGH"),
}

findings: list[tuple[str, str, str]] = []


def add(section: str, severity: str, message: str) -> None:
    findings.append((section, severity, message))


def expand(span: str) -> list[str]:
    """'A-H' or 'A–H' -> [A..H];  'A' -> [A]."""
    span = span.replace("\u2013", "-").replace("\u2014", "-")
    if "-" in span:
        a, b = span.split("-")[:2]
        if a and b and a.isalpha() and b.isalpha():
            return [chr(c) for c in range(ord(a[0]), ord(b[0]) + 1)]
    return [c for c in span if c.isalpha()]


def main() -> None:
    doc = docx.Document(str(DOC))
    paras = [p.text for p in doc.paragraphs]
    try:
        legend_start = next(i for i, t in enumerate(paras) if t.strip() == "Legends")
    except StopIteration:
        legend_start = len(paras)
    try:
        ref_start = next(i for i, t in enumerate(paras) if t.strip() == "References")
    except StopIteration:
        ref_start = legend_start
    body = "\n".join(paras[:ref_start])
    legends = "\n".join(paras[legend_start:])

    # ---- 1. citation order -------------------------------------------------
    def first_mentions(pattern: str, haystack: str) -> list[str]:
        seen: OrderedDict[str, None] = OrderedDict()
        for m in re.finditer(pattern, haystack):
            seen.setdefault(m.group(1), None)
        return list(seen)

    main_order = first_mentions(r"(?<!Supplementary )Fig(?:ure)?\.?\s*(\d+)", body)
    supp_order = first_mentions(r"Supplementary\s+Fig(?:ure)?\.?\s*(\d+)", body)
    tab_order = first_mentions(r"Supplementary\s+Table\s*(\d+)", body)

    for label, order in (
        ("Main figures", main_order),
        ("Supplementary figures", supp_order),
        ("Supplementary tables", tab_order),
    ):
        if order != sorted(order, key=int):
            add(
                "1. Citation order",
                "ERROR",
                f"{label} are not cited in numerical order: first mentions run "
                f"{', '.join(order)}. Journals require numbering by order of first mention.",
            )
        else:
            add("1. Citation order", "OK", f"{label} cited in order: {', '.join(order)}.")

    # ---- 2 & 3. panel letters ---------------------------------------------
    def check_panels(haystack: str, where: str) -> None:
        for m in re.finditer(
            r"(Supplementary\s+)?Fig(?:ure)?\.?\s*(\d+)\s*([A-Z](?:\s*[-\u2013]\s*[A-Z])?)",
            haystack,
        ):
            supp, num, span = bool(m.group(1)), m.group(2), m.group(3)
            table = SUPP_PANELS if supp else FIGURE_PANELS
            name = f"{'Supplementary ' if supp else ''}Figure {num}"
            if num not in table:
                continue
            have = table[num]
            for letter in expand(span):
                if have and letter not in have:
                    add(
                        f"{where}",
                        "ERROR",
                        f"{name}{span} cites panel {letter}, but that figure has panels "
                        f"{have[0]}-{have[-1]} only.",
                    )

    check_panels(body, "2. Panel letters in the text")
    check_panels(legends, "3. Panel letters in the legends")

    # ---- 4. reference order ------------------------------------------------
    sup_numbers: list[int] = []
    for p in doc.paragraphs[:ref_start]:
        for r in p.runs:
            if r.font.superscript and re.fullmatch(r"[\d,\u2013\-\s]+", r.text or ""):
                for chunk in re.split(r"[,\u2013\-]", r.text):
                    if chunk.strip().isdigit():
                        sup_numbers.append(int(chunk))
    first_seen: OrderedDict[int, None] = OrderedDict()
    for n in sup_numbers:
        first_seen.setdefault(n, None)
    cited = list(first_seen)
    n_refs = sum(
        1
        for p in doc.paragraphs[ref_start:legend_start]
        if p.style.name == "List Paragraph" and p.text.strip()
    )
    if cited:
        ascending = cited == sorted(cited)
        add(
            "4. References",
            "OK" if ascending else "ERROR",
            f"{len(cited)} distinct references cited; first-mention order is "
            f"{'ascending' if ascending else 'NOT ascending'}. Reference list holds {n_refs} entries.",
        )
        if not ascending:
            breaks = [
                (cited[i - 1], cited[i]) for i in range(1, len(cited)) if cited[i] < cited[i - 1]
            ]
            add(
                "4. References",
                "ERROR",
                "First-mention sequence breaks at: "
                + "; ".join(f"{a} then {b}" for a, b in breaks[:8])
                + (" ..." if len(breaks) > 8 else "")
                + ". Renumber so citations ascend on first appearance.",
            )
        missing = [n for n in range(1, n_refs + 1) if n not in first_seen]
        if missing:
            add(
                "4. References",
                "WARN",
                f"Reference list entries never cited in the text: {missing}.",
            )
        over = [n for n in cited if n > n_refs]
        if over:
            add(
                "4. References",
                "ERROR",
                f"Text cites reference numbers beyond the list ({n_refs} entries): {over}.",
            )

    # ---- 5. artifact inventory --------------------------------------------
    for num in FIGURE_PANELS:
        f = REPO / "figures" / f"figure{num}.pdf"
        add(
            "5. Repo artifacts",
            "OK" if f.exists() else "ERROR",
            f"figures/figure{num}.pdf {'present' if f.exists() else 'MISSING'}.",
        )
    for num in SUPP_PANELS:
        f = REPO / "figures" / f"supplementary_figure{num}.pdf"
        add(
            "5. Repo artifacts",
            "OK" if f.exists() else "ERROR",
            f"figures/supplementary_figure{num}.pdf {'present' if f.exists() else 'MISSING'}.",
        )

    # ---- write -------------------------------------------------------------
    lines = ["# Manuscript vs repository audit", "", f"Source document: `{DOC.name}`", ""]
    for section in sorted({s for s, _, _ in findings}):
        lines += [f"## {section}", ""]
        for s, sev, msg in findings:
            if s == section:
                mark = {"OK": "OK  ", "WARN": "WARN", "ERROR": "FAIL"}[sev]
                lines.append(f"- `{mark}` {msg}")
        lines.append("")
    OUT.write_text("\n".join(lines), encoding="utf-8")

    n_err = sum(1 for _, s, _ in findings if s == "ERROR")
    n_warn = sum(1 for _, s, _ in findings if s == "WARN")
    print(f"wrote {OUT}")
    print(f"{n_err} errors, {n_warn} warnings, {len(findings)} checks")
    for s, sev, m in findings:
        if sev != "OK":
            print(f"  [{sev}] {s}: {m}")


if __name__ == "__main__":
    main()
