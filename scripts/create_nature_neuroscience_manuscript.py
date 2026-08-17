"""Create the Nature Neuroscience manuscript with precise change highlighting.

The publication-ready revision is compared with the untouched source manuscript.
Only text that is new or changed is highlighted; unchanged text inside revised
paragraphs remains unhighlighted.
"""

from __future__ import annotations

from copy import deepcopy
from difflib import SequenceMatcher
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.run import Run


REPO = Path(__file__).resolve().parents[1]
REPORT = REPO / "report"
REVISED = next(
    path
    for path in REPORT.glob("*publication_ready_highlighted.docx")
    if not path.name.startswith("~$")
)
ORIGINAL = next(
    path
    for path in (Path.home() / "Downloads").glob("Sanguino-G*coping_strategies.docx")
    if path.parent == Path.home() / "Downloads"
)
OUTPUT = REPORT / "Sanguino-Gómez_coping_strategies_Nature_Neuroscience_highlighted.docx"
SI_GUIDE = REPORT / "SIGuide.docx"


ABSTRACT = (
    "Behavioral assays often reduce coping to one response and may miss vulnerability "
    "after early-life stress (ELS). We combined DeepLabCut pose estimation with "
    "supervised SimBA and unsupervised keypoint MoSeq to quantify threat-coping "
    "diversity, organization and predictability in male mice from two ELS cohorts. "
    "ELS offspring showed blunted freezing, prolonged active behaviors and a narrower, "
    "repetitive repertoire. A behavioral-dynamics score identified 12 of 41 ELS mice "
    "(29%) with control-like profiles. These resilient animals showed greater freezing "
    "and Simpson diversity and less recurrence than vulnerable ELS mice. Animal-level "
    "models distinguished resilient from vulnerable animals using full-session motif "
    "frequencies (leave-one-out AUC of 0.842; 95% CI, 0.692–0.951), with variable "
    "held-out-cohort performance (AUC values of 0.817 and 0.632). Exploratory time-resolved "
    "models recovered the full-session label at early horizons, although estimates were "
    "partly dependent because the outcome came from the same session. Computational "
    "ethology therefore revealed flexible, multidimensional coping profiles and promising candidate resilience "
    "markers beyond freezing."
)

RESULTS_VIDEO_SENTENCE = (
    "Representative pose-overlaid examples and canonical skeleton trajectories "
    "for the common syllables are provided in Supplementary Video 1."
)
METHODS_VIDEO_SENTENCE = (
    "For the audiovisual atlas, three 20-frame representative occurrences were "
    "displayed for each available common syllable together with its canonical "
    "14-point skeleton trajectory; the derived climbing class was shown without "
    "a canonical MoSeq skeleton (Supplementary Video 1)."
)
VIDEO_TITLE = (
    "Supplementary Video 1. Representative MoSeq syllables and canonical "
    "skeleton trajectories."
)
VIDEO_LEGEND = (
    "The video presents three representative pose-overlaid occurrences for each "
    "available common syllable (0–34), grouped by the seven curated behavioral "
    "classes or labeled unassigned, alongside the corresponding canonical "
    "14-point skeleton trajectory. Climbing (class 111) is presented separately "
    "because it was defined by the floor-overlap rule rather than as a canonical "
    "MoSeq syllable. The white dot identifies frames assigned to the displayed "
    "syllable. Videos are shown at 25 frames s⁻¹ and contain no audio."
)


TOKEN_RE = re.compile(r"\s+|[\w]+(?:[’'\-–][\w]+)*|[^\w\s]", re.UNICODE)
WORD_RE = re.compile(r"[\w]+(?:[’'\-–][\w]+)*", re.UNICODE)


def paragraph_by_text(document: Document, text: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == text]
    if len(matches) != 1:
        raise AssertionError(f"Expected one paragraph {text!r}, found {len(matches)}")
    return matches[0]


def replace_single_run_text(paragraph, text: str) -> None:
    """Replace text while retaining the paragraph's first-run formatting."""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)


def all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def clear_highlights(document: Document) -> None:
    for paragraph in all_paragraphs(document):
        for run in paragraph.runs:
            run.font.highlight_color = None


def changed_intervals(old_text: str, new_text: str) -> list[tuple[int, int]]:
    """Return character intervals in new_text that differ from old_text."""
    old_tokens = [(match.group(), match.start(), match.end()) for match in TOKEN_RE.finditer(old_text)]
    new_tokens = [(match.group(), match.start(), match.end()) for match in TOKEN_RE.finditer(new_text)]
    matcher = SequenceMatcher(
        None,
        [token[0] for token in old_tokens],
        [token[0] for token in new_tokens],
        autojunk=False,
    )
    intervals: list[tuple[int, int]] = []
    for tag, _i1, _i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal" or j1 == j2:
            continue
        intervals.append((new_tokens[j1][1], new_tokens[j2 - 1][2]))

    merged: list[list[int]] = []
    for start, end in intervals:
        if merged and start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return [(start, end) for start, end in merged]


def is_highlighted(start: int, end: int, intervals: list[tuple[int, int]]) -> bool:
    return any(start < interval_end and end > interval_start for interval_start, interval_end in intervals)


def apply_intervals(paragraph, intervals: list[tuple[int, int]]) -> None:
    """Split runs only where needed, preserving their original character styling."""
    if not intervals or not paragraph.text:
        return

    source_runs = list(paragraph.runs)
    if not source_runs:
        source_runs = [paragraph.add_run(paragraph.text)]

    run_records: list[tuple[Run, int, int]] = []
    cursor = 0
    for run in source_runs:
        start = cursor
        cursor += len(run.text)
        run_records.append((run, start, cursor))

    for run in source_runs:
        run._element.getparent().remove(run._element)

    for source_run, run_start, run_end in run_records:
        if run_start == run_end:
            paragraph._p.append(deepcopy(source_run._r))
            continue
        boundaries = {run_start, run_end}
        for interval_start, interval_end in intervals:
            if run_start < interval_start < run_end:
                boundaries.add(interval_start)
            if run_start < interval_end < run_end:
                boundaries.add(interval_end)
        ordered = sorted(boundaries)
        for start, end in zip(ordered, ordered[1:]):
            new_r_element = deepcopy(source_run._r)
            new_run = Run(new_r_element, paragraph)
            new_run.text = source_run.text[start - run_start : end - run_start]
            new_run.font.highlight_color = (
                WD_COLOR_INDEX.YELLOW if is_highlighted(start, end, intervals) else None
            )
            paragraph._p.append(new_r_element)


def mark_precise_changes(original: Document, revised: Document) -> None:
    old_paragraphs = original.paragraphs
    new_paragraphs = revised.paragraphs
    matcher = SequenceMatcher(
        None,
        [paragraph.text for paragraph in old_paragraphs],
        [paragraph.text for paragraph in new_paragraphs],
        autojunk=False,
    )

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal" or tag == "delete":
            continue
        if tag == "insert":
            for paragraph in new_paragraphs[j1:j2]:
                apply_intervals(paragraph, [(0, len(paragraph.text))])
            continue

        paired = min(i2 - i1, j2 - j1)
        for offset in range(paired):
            old_text = old_paragraphs[i1 + offset].text
            paragraph = new_paragraphs[j1 + offset]
            apply_intervals(paragraph, changed_intervals(old_text, paragraph.text))
        for paragraph in new_paragraphs[j1 + paired : j2]:
            apply_intervals(paragraph, [(0, len(paragraph.text))])


def append_sentence(paragraph, sentence: str) -> None:
    separator = "" if paragraph.text.endswith((" ", "\t", "\n")) else " "
    paragraph.add_run(separator + sentence)


def write_si_guide() -> None:
    guide = Document()
    guide.add_heading("Supplementary Information Guide", level=1)
    guide.add_paragraph(
        "Coping dynamics reveal individual resilience following early-life stress"
    )
    title = guide.add_paragraph()
    title.add_run(VIDEO_TITLE).bold = True
    guide.add_paragraph(VIDEO_LEGEND)
    guide.save(SI_GUIDE)


def main() -> None:
    original = Document(ORIGINAL)
    revised = Document(REVISED)

    abstract_index, abstract_heading = next(
        (index, paragraph)
        for index, paragraph in enumerate(revised.paragraphs)
        if paragraph.text.startswith("Abstract (")
    )
    replace_single_run_text(abstract_heading, "Abstract (150 words)")
    replace_single_run_text(revised.paragraphs[abstract_index + 1], ABSTRACT)

    methods_heading = paragraph_by_text(revised, "Methods")
    replace_single_run_text(methods_heading, "Online Methods")

    results_paragraph = next(
        paragraph
        for paragraph in revised.paragraphs
        if paragraph.text.startswith("Once validated, we used keypoint MoSeq")
    )
    append_sentence(results_paragraph, RESULTS_VIDEO_SENTENCE)

    methods_paragraph = next(
        paragraph
        for paragraph in revised.paragraphs
        if paragraph.text.startswith(
            "To interpret the behavioral relevance of the keypoint MoSeq syllables"
        )
    )
    append_sentence(methods_paragraph, METHODS_VIDEO_SENTENCE)

    revised.add_heading("Supplementary Videos", level=1)
    title_paragraph = revised.add_paragraph()
    title_paragraph.add_run(VIDEO_TITLE).bold = True
    revised.add_paragraph(VIDEO_LEGEND)

    clear_highlights(revised)
    mark_precise_changes(original, revised)
    revised.save(OUTPUT)
    write_si_guide()

    abstract_words = len(ABSTRACT.split())
    highlighted_runs = sum(
        run.font.highlight_color == WD_COLOR_INDEX.YELLOW
        for paragraph in all_paragraphs(revised)
        for run in paragraph.runs
    )
    print(f"source: {ORIGINAL}")
    print(f"revised base: {REVISED}")
    print(f"abstract word count: {abstract_words}")
    print(f"video legend word count: {len(VIDEO_LEGEND.split())}")
    print(f"highlighted runs: {highlighted_runs}")
    print(f"wrote: {OUTPUT}")
    print(f"wrote: {SI_GUIDE}")


if __name__ == "__main__":
    main()
