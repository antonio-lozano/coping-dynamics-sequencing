# SPDX-License-Identifier: MIT
"""Build a final Excel report styled like STATISTICAL_REPORT_FINAL.xlsx.

The project already writes analysis outputs as CSV files and simple workbooks.
This exporter keeps those outputs as the source of truth, but borrows the
visual grammar from a hand-formatted reference workbook: merged sheet titles,
Calibri report text, styled headers, restrained borders, and red p-value
significance highlighting.
"""
from __future__ import annotations

import argparse
import math
import re
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = Path.home() / "Downloads" / "STATISTICAL_REPORT_FINAL.xlsx"
DEFAULT_OUTPUT = REPO_ROOT / "results" / "statistical_reports" / "STATISTICAL_REPORT_FINAL_rebuilt.xlsx"

CSV_PATTERNS = (
    "results/statistical_reports/*.csv",
    "results/figure_data/stats_*.csv",
    "results/figure_data/source_data_*.csv",
    "results/figure_data/*threshold_audit.csv",
    "results/figure_data/supplementary_figure*_*.csv",
    "results/figure_data/transition_metrics_per_animal.csv",
)

P_VALUE_RE = re.compile(r"(^p$|p[_ .-]?(value|display|adj)|fdr|q[_ .-]?value|pr[.(]|p>|bh)", re.I)
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _clone_style(src: Cell, dst: Cell) -> None:
    if src.has_style:
        dst._style = copy(src._style)
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def _fallback_styles() -> dict[str, Cell]:
    wb = Workbook()
    ws = wb.active

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"].font = Font(name="Calibri", bold=True, size=12, color="4D4D4D")
    ws["A1"].fill = PatternFill("solid", fgColor="D9EAD3")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A2"].font = Font(name="Calibri", size=11, color="4D4D4D")

    ws["A3"].font = Font(name="Calibri", bold=True, size=11, color="4D4D4D")
    ws["A3"].fill = PatternFill("solid", fgColor="EAF2F8")
    ws["A3"].border = border
    ws["A3"].alignment = Alignment(wrap_text=True)

    ws["A4"].font = Font(name="Calibri", size=11, color="4D4D4D")
    ws["A4"].border = border
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")

    ws["B4"].font = Font(name="Calibri", size=11, color="4D4D4D")
    ws["B4"].border = border
    ws["B4"].number_format = "0.000000"
    ws["B4"].alignment = Alignment(vertical="top")

    ws["C4"].font = Font(name="Calibri", size=11, color="C0392B")
    ws["C4"].border = border
    ws["C4"].number_format = "0.000000"
    ws["C4"].alignment = Alignment(vertical="top")

    return {
        "title": ws["A1"],
        "blank": ws["A2"],
        "header": ws["A3"],
        "text": ws["A4"],
        "number": ws["B4"],
        "significant": ws["C4"],
    }


def _template_styles(template_path: Path) -> dict[str, Cell]:
    if not template_path.exists():
        return _fallback_styles()

    wb = load_workbook(template_path, read_only=False, data_only=False)

    def cell(sheet: str, address: str, fallback: tuple[str, str]) -> Cell:
        if sheet in wb.sheetnames:
            return wb[sheet][address]
        return wb[fallback[0]][fallback[1]]

    first_sheet = wb.sheetnames[0]
    styles = {
        "title": cell("Fig.2D_Syllable_usage", "A1", (first_sheet, "A1")),
        "blank": cell("Fig.2D_Syllable_usage", "A2", (first_sheet, "A2")),
        "header": cell("Fig.2D_Syllable_usage", "A3", (first_sheet, "A3")),
        "text": cell("Fig.2D_Syllable_usage", "A4", (first_sheet, "A4")),
        "number": cell("Fig.2D_Syllable_usage", "C4", (first_sheet, "A4")),
        "significant": cell("Fig.5B_Euclidean distance", "E4", (first_sheet, "A4")),
    }
    return styles


def _iter_csvs(extra_csvs: Iterable[Path] = ()) -> list[Path]:
    paths: list[Path] = []
    for pattern in CSV_PATTERNS:
        paths.extend(REPO_ROOT.glob(pattern))
    paths.extend(Path(p) for p in extra_csvs)

    unique: dict[Path, Path] = {}
    for path in paths:
        resolved = path if path.is_absolute() else (REPO_ROOT / path)
        if resolved.exists() and resolved.suffix.lower() == ".csv":
            unique[resolved.resolve()] = resolved
    return sorted(unique.values(), key=lambda p: (str(p.parent), p.name.lower()))


def _safe_sheet_name(path: Path, used: set[str]) -> str:
    name = path.stem
    name = name.replace("source_data_", "src_")
    name = name.replace("supplementary_", "supp_")
    name = name.replace("statistical_", "stats_")
    name = INVALID_SHEET_CHARS.sub("_", name)
    name = re.sub(r"\s+", "_", name).strip("_") or "sheet"
    base = name[:31]
    candidate = base
    i = 2
    while candidate in used:
        suffix = f"_{i}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        i += 1
    used.add(candidate)
    return candidate


def _is_p_value_column(column_name: object) -> bool:
    return bool(P_VALUE_RE.search(str(column_name)))


def _is_significant(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return False
    if isinstance(value, float) and math.isnan(value):
        return False
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return False
        if text.startswith("<"):
            try:
                return float(text[1:]) <= 0.05
            except ValueError:
                return False
        text = text.replace(",", ".")
        try:
            return float(text) < 0.05
        except ValueError:
            return False
    try:
        return float(value) < 0.05
    except (TypeError, ValueError):
        return False


def _mark_significant(cell: Cell) -> None:
    font = copy(cell.font)
    font.color = "C0392B"
    cell.font = font


def _excel_value(value: object) -> object:
    if pd.isna(value):
        return None
    return value


def _content_width(values: Iterable[object], header: object) -> float:
    lengths = [len(str(header))]
    for value in values:
        if value is not None and not (isinstance(value, float) and math.isnan(value)):
            lengths.append(len(str(value)))
    return min(max(max(lengths) + 2, 8), 55)


def _write_dataframe_sheet(wb: Workbook, styles: dict[str, Cell], sheet_name: str, title: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title=sheet_name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

    if df.empty:
        df = pd.DataFrame([{"note": "No rows in source file"}])

    n_cols = max(len(df.columns), 1)
    last_col = get_column_letter(n_cols)

    ws.cell(1, 1, title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    for col in range(1, n_cols + 1):
        _clone_style(styles["title"], ws.cell(1, col))
    ws.row_dimensions[1].height = 24

    for col in range(1, n_cols + 1):
        _clone_style(styles["blank"], ws.cell(2, col))

    for col, header in enumerate(df.columns, start=1):
        cell = ws.cell(3, col, str(header))
        _clone_style(styles["header"], cell)
    ws.auto_filter.ref = f"A3:{last_col}{len(df) + 3}"

    p_value_cols = {_idx for _idx, name in enumerate(df.columns, start=1) if _is_p_value_column(name)}
    for row_idx, row in enumerate(df.itertuples(index=False, name=None), start=4):
        for col_idx, value in enumerate(row, start=1):
            out = _excel_value(value)
            cell = ws.cell(row_idx, col_idx, out)
            style_key = "number" if isinstance(out, (int, float)) and not isinstance(out, bool) else "text"
            if col_idx in p_value_cols and _is_significant(out):
                style_key = "significant"
            _clone_style(styles[style_key], cell)
            if style_key == "significant":
                _mark_significant(cell)

    for col_idx, header in enumerate(df.columns, start=1):
        values = df.iloc[:, col_idx - 1].head(300).tolist()
        ws.column_dimensions[get_column_letter(col_idx)].width = _content_width(values, header)


def _write_index_sheet(wb: Workbook, styles: dict[str, Cell], csv_paths: list[Path], template_path: Path) -> None:
    rows = [
        {"field": "Generated", "value": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        {"field": "Template", "value": str(template_path)},
        {"field": "Source root", "value": str(REPO_ROOT)},
        {"field": "Included CSV files", "value": len(csv_paths)},
    ]
    rows.extend({"field": f"CSV {i}", "value": str(path.relative_to(REPO_ROOT))} for i, path in enumerate(csv_paths, start=1))
    _write_dataframe_sheet(wb, styles, "Report_index", "Final report index", pd.DataFrame(rows))


def build_report(template_path: Path, output_path: Path, extra_csvs: Iterable[Path] = ()) -> Path:
    styles = _template_styles(template_path)
    csv_paths = _iter_csvs(extra_csvs)
    if not csv_paths:
        raise FileNotFoundError("No CSV outputs were found to include in the final report.")

    wb = Workbook()
    wb.remove(wb.active)

    _write_index_sheet(wb, styles, csv_paths, template_path)

    used = {"Report_index"}
    for csv_path in csv_paths:
        df = pd.read_csv(csv_path)
        sheet_name = _safe_sheet_name(csv_path, used)
        title = csv_path.stem.replace("_", " ")
        _write_dataframe_sheet(wb, styles, sheet_name, title, df)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE, help="Reference workbook to borrow styles from.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Styled final workbook to write.")
    parser.add_argument("--csv", type=Path, action="append", default=[], help="Additional CSV file to include; may be repeated.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output = build_report(args.template, args.output, args.csv)
    print(f"Saved styled final report: {output}")


if __name__ == "__main__":
    main()
