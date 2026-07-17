"""Check whether Figure 4 Report.xlsx values could come from n=81.

For each Fig. 4 model, rerun the old-code Control-vs-ELS MixedLM after dropping
one animal at a time, then rank the n=81 fits by closeness to Report.xlsx.
"""
from __future__ import annotations

import warnings
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import entropy
import statsmodels.formula.api as smf

warnings.filterwarnings("ignore")

REPO = Path(__file__).resolve().parents[1]
RAW_250MS = REPO / "data" / "source" / "syllable_usage_per_timebin_250ms.csv"
REPORT_XLSX = Path("E:/#Deeplabcut_project/Report/Report.xlsx")

CLUSTER_MAP = {
    "Freezing": [0, 28],
    "Sniffing": [18, 20],
    "Grooming": [24],
    "Turn": [1, 3, 5, 6, 10, 15, 26, 27],
    "Locomotion": [11, 12, 14, 16, 19, 21, 25],
    "Climbing": [111],
    "Jump": [23, 29, 30, 34],
}


def load_raw(drop_animal: str | None = None) -> pd.DataFrame:
    raw = pd.read_csv(RAW_250MS).rename(columns={"Time Bin": "Time_bin"})
    raw = raw[raw["Condition"].isin(["Control", "ELS"])].copy()
    raw["Animal"] = raw["Animal"].astype(str)
    if drop_animal is not None:
        raw = raw[raw["Animal"] != drop_animal].copy()
    raw["Syllable"] = pd.to_numeric(raw["Syllable"], errors="coerce").astype(int)
    raw["Experiment"] = pd.to_numeric(raw["Experiment"], errors="coerce")
    s2c = {int(s): c for c, ss in CLUSTER_MAP.items() for s in ss}
    raw["Cluster"] = raw["Syllable"].map(s2c).fillna("")
    return raw


def frequency_metrics(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    freq_df = raw.groupby(["Animal", "Experiment", "Condition", "Cluster"]).size().reset_index(name="Count")
    pivot = freq_df.pivot_table(
        index=["Animal", "Condition", "Experiment"],
        columns="Cluster",
        values="Count",
        fill_value=0,
    )
    rel = pivot.div(pivot.sum(axis=1), axis=0)
    ent = rel.apply(lambda row: entropy(row[row > 0]), axis=1)
    nclus = (pivot > 0).sum(axis=1)
    even = ent / np.log(nclus)
    simp = rel.apply(lambda row: 1 - np.sum(row**2), axis=1)
    metrics = pd.DataFrame({"Shannon": ent, "Evenness": even, "Simpson": simp}).reset_index()
    metrics["group"] = metrics["Condition"]
    return metrics, freq_df


def cui_df(freq_df: pd.DataFrame, raw: pd.DataFrame) -> pd.DataFrame:
    freq_use = freq_df.drop(columns=["Condition"]).copy()
    merged = freq_use.merge(raw[["Animal", "Condition"]].drop_duplicates(), on="Animal", how="left")
    total = merged.groupby(["Animal", "Condition", "Experiment"])["Count"].sum().reset_index(name="Total")
    usage = merged.merge(total, on=["Animal", "Condition", "Experiment"])
    usage["RelativeUsage"] = usage["Count"] / usage["Total"]
    usage = usage[usage["Cluster"].astype(str).str.strip() != ""].copy()

    def compute(subdf: pd.DataFrame) -> float:
        cum = subdf.sort_values("RelativeUsage", ascending=False)["RelativeUsage"].cumsum().to_numpy()
        n = len(cum)
        baseline = (n + 1) / (2 * n)
        return float((cum.mean() - baseline) / (1 - baseline))

    out = usage.groupby(["Animal", "Experiment", "Condition"]).apply(compute).reset_index(name="CUI")
    out["group"] = out["Condition"]
    return out


def bouts_df(raw: pd.DataFrame) -> pd.DataFrame:
    def compute(subdf: pd.DataFrame) -> pd.Series:
        subdf = subdf.sort_values("Time_bin").copy()
        subdf["Cluster_change"] = (subdf["Cluster"] != subdf["Cluster"].shift(1)).astype(int)
        subdf["Bout"] = subdf["Cluster_change"].cumsum()
        bouts = subdf.groupby(["Bout", "Cluster"]).size() * 0.25
        return bouts.groupby(level=1).mean()

    out = raw.groupby(["Animal", "Experiment"]).apply(compute).reset_index()
    out.columns = ["Animal", "Experiment", "Cluster", "MeanBoutDuration"]
    out = out.dropna(subset=["MeanBoutDuration"]).reset_index(drop=True)
    out = out.merge(raw[["Animal", "Condition", "Experiment"]].drop_duplicates(), on=["Animal", "Experiment"], how="left")
    out["group"] = out["Condition"]
    return out


def fit(df: pd.DataFrame, y: str) -> tuple[float, float, float, float, int]:
    sub = df.dropna(subset=[y]).copy()
    sub["group"] = pd.Categorical(sub["group"], categories=["Control", "ELS"])
    model = smf.mixedlm(f"{y} ~ group + Experiment", sub, groups=sub["Animal"]).fit(reml=False)
    key = "group[T.ELS]"
    return float(model.params[key]), float(model.bse[key]), float(model.tvalues[key]), float(model.pvalues[key]), int(model.nobs)


def report_value(sheet: str, start_col: int) -> tuple[float, float, float, float, int]:
    wb = load_workbook(REPORT_XLSX, read_only=True, data_only=True)
    ws = wb[sheet]
    return (
        float(ws.cell(4, start_col + 1).value),
        float(ws.cell(4, start_col + 2).value),
        float(ws.cell(4, start_col + 3).value),
        float(ws.cell(4, start_col + 4).value),
        int(ws.cell(6, start_col + 10).value),
    )


def score(vals: tuple[float, float, float, float, int], rep: tuple[float, float, float, float, int]) -> float:
    # Scale beta/se/z differences by report magnitudes; p is noisier, so lower weight.
    b, se, z, p, _ = vals
    rb, rse, rz, rp, _ = rep
    return (
        abs(b - rb) / max(abs(rb), 1e-9)
        + abs(se - rse) / max(abs(rse), 1e-9)
        + abs(z - rz) / max(abs(rz), 1e-9)
        + 0.25 * abs(p - rp) / max(abs(rp), 1e-9)
    )


def fmt(vals: tuple[float, float, float, float, int]) -> str:
    return f"b={vals[0]:+.6f} se={vals[1]:.6f} z={vals[2]:+.6f} p={vals[3]:.6g} n={vals[4]}"


def build_tables_from_raw(raw: pd.DataFrame) -> dict[str, tuple[pd.DataFrame, str]]:
    metrics, freq = frequency_metrics(raw)
    cui = cui_df(freq, raw)
    bouts = bouts_df(raw)
    return {
        "Simpson": (metrics, "Simpson"),
        "CUI": (cui, "CUI"),
        "Freezing bout": (bouts[bouts["Cluster"] == "Freezing"], "MeanBoutDuration"),
        "Sniffing bout": (bouts[bouts["Cluster"] == "Sniffing"], "MeanBoutDuration"),
        "Turn bout": (bouts[bouts["Cluster"] == "Turn"], "MeanBoutDuration"),
    }


def drop_from_tables(
    tables: dict[str, tuple[pd.DataFrame, str]], animal: str | None
) -> dict[str, tuple[pd.DataFrame, str]]:
    if animal is None:
        return tables
    out: dict[str, tuple[pd.DataFrame, str]] = {}
    for label, (df, y) in tables.items():
        out[label] = (df[df["Animal"] != animal].copy(), y)
    return out


def main() -> None:
    report = {
        "Simpson": report_value("Frequency_metrics", 1),
        "CUI": report_value("Frequency_metrics", 41),
        "Freezing bout": report_value("Bout_duration", 14),
        "Sniffing bout": report_value("Bout_duration", 27),
        "Turn bout": report_value("Bout_duration", 54),
    }
    raw = load_raw()
    animals = sorted(raw["Animal"].unique())
    full_tables = build_tables_from_raw(raw)
    print("Report workbook N metadata")
    for label, rep in report.items():
        full = fit(*full_tables[label])
        print(f"{label:14s} report {fmt(rep)} | n=82 old-code {fmt(full)}")

    print("\nBest n=81 leave-one-out candidates")
    for label, rep in report.items():
        candidates = []
        for animal in animals:
            try:
                vals = fit(*drop_from_tables(full_tables, animal)[label])
            except Exception:
                continue
            candidates.append((score(vals, rep), animal, vals))
        candidates.sort(key=lambda x: x[0])
        print(f"\n{label} report: {fmt(rep)}")
        for sc, animal, vals in candidates[:8]:
            print(f"  drop {animal:>6s} score={sc:.4f} {fmt(vals)}")


if __name__ == "__main__":
    main()
