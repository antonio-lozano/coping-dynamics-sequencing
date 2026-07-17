"""Verify the manuscript freezing Results subsection against Report.xlsx.

Checks:
- Figure 1 SimBA validation correlation from Raw_data.xlsx.
- Figure 2A-C freezing MixedLM values from raw prediction CSVs.
- External lab report values in E:/#Deeplabcut_project/Report/Report.xlsx.
"""

from __future__ import annotations

import re
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import statsmodels.formula.api as smf
from openpyxl import load_workbook
from scipy import stats


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "source"
FREEZING_DIR = SOURCE / "freezing_predictions"
REPORT_PATH = Path(r"E:\#Deeplabcut_project\Report\Report.xlsx")
RAW_WORKBOOK = ROOT / "results" / "figure_data" / "Raw_data.xlsx"

FPS = 25
BIN_SECONDS = 30
EXCLUDE = {"animal_48_6", "48_6"}

MANUSCRIPT = {
    "Fig1 validation": {"r": 0.95, "r2": 0.90, "p_label": "<0.001"},
    "Exp1 time": {"beta": 3.721, "se": 0.128, "z": 29.048, "p_label": "<0.001"},
    "Exp1 ELS x time": {"beta": -0.887, "se": 0.181, "z": -4.898, "p_label": "<0.001"},
    "Exp3 time": {"beta": 4.412, "se": 0.191, "z": 23.091, "p_label": "<0.001"},
    "Exp3 ELS x time": {"beta": -0.721, "se": 0.270, "z": -2.667, "p_label": "0.008"},
    "Combined ELS x time": {"beta": -0.822, "se": 0.154, "z": -5.327, "p_label": "<0.001"},
}


def normalize_name(value: str) -> str:
    return str(value).strip().replace(" ", "_")


def short_id(value: str) -> str:
    norm = normalize_name(value)
    match = re.search(r"Animal[_ ]?(\d+(?:[_-]\d+)?)", norm, re.IGNORECASE)
    if match:
        return match.group(1).replace("_", ".").replace("-", ".")
    parts = norm.split("_")
    return ".".join(parts[-2:]) if len(parts) >= 2 else norm


def fmt_p(value: float) -> str:
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def report_match(label: str, manuscript: dict[str, float | str], observed: dict[str, float]) -> None:
    pieces = []
    for key in ("beta", "se", "z"):
        diff = abs(float(manuscript[key]) - observed[key])
        tag = "MATCH" if diff <= 0.001 else ("CLOSE" if diff <= 0.01 else "DIFF")
        pieces.append(f"{key}: text={float(manuscript[key]):.3f} observed={observed[key]:.3f} {tag}")
    pieces.append(f"p: text={manuscript['p_label']} observed={fmt_p(observed['p'])}")
    print(f"{label}: " + "; ".join(pieces))


def load_freezing_long() -> pd.DataFrame:
    idx = pd.read_csv(SOURCE / "animal_groups.csv")
    group_map: dict[str, str] = {}
    for _, row in idx.iterrows():
        raw = str(row["name"]).strip()
        group = str(row["group"])
        for key in {raw, normalize_name(raw), short_id(raw)}:
            group_map[key] = group

    freq = pd.read_csv(SOURCE / "cluster_frequency_per_animal.csv")
    exp_map = dict(zip(freq["animal_id"].astype(str), freq["experiment"]))

    records: list[dict[str, float | str]] = []
    bin_size = FPS * BIN_SECONDS
    for csv_path in sorted(FREEZING_DIR.glob("*_freezing_predictions_only.csv")):
        base = csv_path.name.replace("_freezing_predictions_only.csv", "")
        if normalize_name(base).lower() in EXCLUDE:
            continue
        df = pd.read_csv(csv_path)
        if "Freezing_Jen_0-125_threshold" not in df.columns:
            continue
        sid = short_id(base)
        group = group_map.get(base) or group_map.get(normalize_name(base)) or group_map.get(sid)
        if group not in ("Control", "ELS"):
            continue
        values = df["Freezing_Jen_0-125_threshold"].to_numpy(dtype=float)
        for bin_idx in range(len(values) // bin_size):
            chunk = values[bin_idx * bin_size : (bin_idx + 1) * bin_size]
            records.append(
                {
                    "animal_id": sid,
                    "group": group,
                    "time_bin": bin_idx + 1,
                    "freezing_pct": float(np.mean(chunk) * 100.0),
                }
            )

    long_df = pd.DataFrame(records)

    def exp_for_sid(sid: str) -> float | None:
        try:
            key = str(float(str(sid).replace("_", ".")))
        except ValueError:
            key = str(sid)
        return exp_map.get(key, exp_map.get(str(sid)))

    long_df["experiment"] = long_df["animal_id"].map(exp_for_sid)
    long_df = long_df.dropna(subset=["experiment"]).copy()
    long_df["experiment"] = long_df["experiment"].astype(int)
    long_df["condition"] = pd.Categorical(long_df["group"], categories=["Control", "ELS"])
    return long_df


def fit_model(sub: pd.DataFrame, use_experiment_vc: bool = False):
    formula = "freezing_pct ~ condition * time_bin"
    if use_experiment_vc:
        return smf.mixedlm(
            formula,
            sub,
            groups=sub["animal_id"],
            re_formula="1",
            vc_formula={"experiment": "0 + C(experiment)"},
        ).fit(reml=False, method="lbfgs", maxiter=200)
    return smf.mixedlm(formula, sub, groups=sub["animal_id"], re_formula="1").fit(
        reml=False, method="lbfgs", maxiter=200
    )


def term(model, name: str) -> dict[str, float]:
    return {
        "beta": float(model.params[name]),
        "se": float(model.bse[name]),
        "z": float(model.tvalues[name]),
        "p": float(model.pvalues[name]),
    }


def read_report_ground_truth() -> dict[str, dict[str, float]]:
    wb = load_workbook(REPORT_PATH, read_only=True, data_only=True)
    ws = wb["Ground_truth "]
    sections = {
        "Combined": 1,
        "Exp1": 19,
        "Exp3": 37,
    }
    rows = {"time": 5, "interaction": 6}
    out: dict[str, dict[str, float]] = {}
    for section, start_col in sections.items():
        for row_label, row in rows.items():
            key = f"{section} {row_label}"
            out[key] = {
                "beta": float(ws.cell(row, start_col + 1).value),
                "se": float(ws.cell(row, start_col + 2).value),
                "z": float(ws.cell(row, start_col + 3).value),
                "p": float(ws.cell(row, start_col + 4).value),
            }
    return out


def check_simba_validation() -> None:
    df = pd.read_excel(RAW_WORKBOOK, sheet_name="SimBA_validation", header=None)
    numeric = df.apply(pd.to_numeric, errors="coerce")
    candidates = [col for col in numeric.columns if numeric[col].notna().sum() > 1000]
    print("\nFigure 1 SimBA validation candidates")
    if len(candidates) < 2:
        print("Could not identify paired validation columns in SimBA_validation.")
        return
    for idx, left in enumerate(candidates):
        for right in candidates[idx + 1 :]:
            sub = numeric[[left, right]].dropna()
            if len(sub) <= 1000:
                continue
            r_value, p_value = stats.pearsonr(sub[left], sub[right])
            r2 = r_value**2
            print(
                f"cols {left},{right}: n={len(sub)} r={r_value:.6f} "
                f"R2={r2:.6f} p={p_value:.3g}"
            )


def main() -> None:
    warnings.filterwarnings("ignore")

    print(f"Report workbook: {REPORT_PATH}")
    report = read_report_ground_truth()
    print("\nExternal Report.xlsx Ground_truth values")
    for key in ["Exp1 time", "Exp1 interaction", "Exp3 time", "Exp3 interaction", "Combined interaction"]:
        print(f"{key}: {report[key]}")

    long_df = load_freezing_long()
    print(
        "\nRaw freezing data: "
        f"{long_df['animal_id'].nunique()} animals, {len(long_df)} observations; "
        f"Exp1={long_df[long_df['experiment'] == 1]['animal_id'].nunique()} animals, "
        f"Exp3={long_df[long_df['experiment'] == 3]['animal_id'].nunique()} animals"
    )

    exp1 = fit_model(long_df[long_df["experiment"] == 1].copy())
    exp3 = fit_model(long_df[long_df["experiment"] == 3].copy())
    combined = fit_model(long_df.copy(), use_experiment_vc=True)

    model_values = {
        "Exp1 time": term(exp1, "time_bin"),
        "Exp1 ELS x time": term(exp1, "condition[T.ELS]:time_bin"),
        "Exp3 time": term(exp3, "time_bin"),
        "Exp3 ELS x time": term(exp3, "condition[T.ELS]:time_bin"),
        "Combined time": term(combined, "time_bin"),
        "Combined ELS x time": term(combined, "condition[T.ELS]:time_bin"),
    }

    print("\nRerun MixedLM values from raw predictions")
    for key, value in model_values.items():
        print(f"{key}: {value}")

    print("\nManuscript text vs rerun model")
    for key in ["Exp1 time", "Exp1 ELS x time", "Exp3 time", "Exp3 ELS x time", "Combined ELS x time"]:
        report_match(key, MANUSCRIPT[key], model_values[key])

    check_simba_validation()


if __name__ == "__main__":
    main()
