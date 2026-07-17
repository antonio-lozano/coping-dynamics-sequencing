import pandas as pd, numpy as np
from scipy import stats
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ── Load and filter to Context only ───────────────────────────────────────
df = pd.read_excel(r'C:\Users\jenif\Downloads\Copying_dynamics_paper2\Raw_data (1).xlsx',
                   sheet_name='SimBA_validation', header=None)
df.columns = df.iloc[1]
df = df.iloc[2:].reset_index(drop=True)
df.columns = ['phase', 'animal_index', 'manual_percent', 'automatic_percent']
df['manual_percent'] = pd.to_numeric(df['manual_percent'], errors='coerce')
df['automatic_percent'] = pd.to_numeric(df['automatic_percent'], errors='coerce')
df = df[df['phase'] == 'Context'].dropna(subset=['manual_percent', 'automatic_percent']).reset_index(drop=True)

print(f"Context only: {len(df)} rows, {df['animal_index'].nunique()} animals")

# ── Compute statistics ─────────────────────────────────────────────────────
x = df['manual_percent'].values
y = df['automatic_percent'].values

r, p = stats.pearsonr(x, y)
slope, intercept, _, _, se = stats.linregress(x, y)

result = {
    'phase': 'Context', 'n': len(x),
    'r': round(r, 4), 'r_squared': round(r**2, 4),
    'p_value': float(p), 'slope': round(slope, 4),
    'intercept': round(intercept, 4), 'SE_slope': round(se, 4),
}
print("Stats:", result)

# ── Styling ────────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill('solid', fgColor='2F5496')
HEADER_FONT = Font(bold=True, color='FFFFFF', size=10)
TITLE_FILL  = PatternFill('solid', fgColor='D6E4F7')
TITLE_FONT  = Font(bold=True, size=11)
thin = Side(style='thin', color='AAAAAA')
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)

def write_title(ws, title, n_cols):
    ws.cell(row=1, column=1, value=title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(row=1, column=1)
    c.fill = TITLE_FILL; c.font = TITLE_FONT
    c.alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = 18

def write_header(ws, row, cols):
    for ci, col in enumerate(cols, 1):
        c = ws.cell(row=row, column=ci, value=col)
        c.fill = HEADER_FILL; c.font = HEADER_FONT
        c.alignment = Alignment(horizontal='center'); c.border = BORD

# ── 1. Raw_data.xlsx — SimBA_validation: Context data only ────────────────
RAW_PATH = r'D:\coping-dynamics-sequencing\results\figure_data\Raw_data.xlsx'
wb = load_workbook(RAW_PATH)
ws = wb['SimBA_validation']
ws.delete_rows(1, ws.max_row)

write_title(ws, 'SimBA_validation - Context phase: Manual vs Automatic Freezing % (n=47 animals)', 4)
write_header(ws, 2, ['phase', 'animal_index', 'manual_percent', 'automatic_percent'])
for ri, rec in enumerate(df.itertuples(index=False), start=3):
    ws.cell(row=ri, column=1, value=rec.phase)
    ws.cell(row=ri, column=2, value=rec.animal_index)
    ws.cell(row=ri, column=3, value=round(float(rec.manual_percent), 4))
    ws.cell(row=ri, column=4, value=round(float(rec.automatic_percent), 4))

for ci, w in zip(range(1, 5), [12, 16, 18, 22]):
    ws.column_dimensions[get_column_letter(ci)].width = w

wb.save(RAW_PATH)
print('Raw_data.xlsx updated')

# ── 2. Statistical_report.xlsx — Fig.1A: Context stats only ───────────────
STAT_PATH = r'D:\coping-dynamics-sequencing\results\figure_data\Statistical_report.xlsx'
wb2 = load_workbook(STAT_PATH)
ws2 = wb2['Fig.1A_ SimBA_validation']
ws2.delete_rows(1, ws2.max_row)

cols = ['phase', 'n', 'r', 'r_squared', 'p_value', 'slope', 'intercept', 'SE_slope', 'significance']
write_title(ws2, 'Fig. 1A - SimBA Validation: Context phase (n=47 animals)', len(cols))
write_header(ws2, 2, cols)

ws2.cell(row=3, column=1, value=result['phase'])
ws2.cell(row=3, column=2, value=result['n'])
ws2.cell(row=3, column=3, value=result['r'])
ws2.cell(row=3, column=4, value=result['r_squared'])
ws2.cell(row=3, column=5, value=result['p_value'])
ws2.cell(row=3, column=6, value=result['slope'])
ws2.cell(row=3, column=7, value=result['intercept'])
ws2.cell(row=3, column=8, value=result['SE_slope'])
ws2.cell(row=3, column=9, value='p < 0.001' if result['p_value'] < 0.001 else 'p = %.4f' % result['p_value'])

for ci, w in zip(range(1, 10), [12, 6, 8, 12, 16, 10, 12, 12, 14]):
    ws2.column_dimensions[get_column_letter(ci)].width = w

wb2.save(STAT_PATH)
print('Statistical_report.xlsx updated')
